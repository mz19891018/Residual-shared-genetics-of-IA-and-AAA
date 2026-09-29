#!/usr/bin/env python
# -*- coding: utf-8 -*-
from docx import Document

# Check v3_final for user's replacement text
doc = Document(r'D:\IA_AAA_analysis\Manuscript_v3_final.docx')

print('=== Searching for replacement text markers in v3_final ===')
for i, para in enumerate(doc.paragraphs):
    text = para.text
    # Look for common markers of user-added replacement text
    if any(kw in text for kw in ['REPLACE', 'replace', 'INSERT', 'insert', 'NEW TEXT', 'new text', '替换', '粘贴', 'paste here', 'use this', 'Use this']):
        print(f'[{i}] ({para.style.name}) {text[:300]}')
        print()

# Also check for eQTL and FUSION method sections
print('=== Methods sections in v3_final ===')
for i, para in enumerate(doc.paragraphs):
    if para.style.name == 'Heading 3' and any(kw in para.text for kw in ['eQTL', 'FUSION', 'TWAS', 'colocaliz', 'Mendelian']):
        print(f'[{i}] HEADING: {para.text}')
        # Print next 3 paragraphs
        for j in range(i+1, min(i+4, len(doc.paragraphs))):
            if doc.paragraphs[j].text.strip():
                print(f'  [{j}] {doc.paragraphs[j].text[:200]}')
        print()
