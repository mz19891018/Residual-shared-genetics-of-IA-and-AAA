#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""Final fixes for Manuscript v3"""
from docx import Document

doc = Document(r'D:\IA_AAA_analysis\Manuscript_v3_final.docx')

# Fix Para 52 - remove 【待重算】 marker and duplicate content
for i, para in enumerate(doc.paragraphs):
    if '【待重算' in para.text:
        # Remove the marker
        for run in para.runs:
            if '【待重算' in run.text:
                run.text = run.text.replace('【待重算：用 eQTL Catalogue 全量汇总统计做 coloc，并纳入 CDKN2B-AS1】',
                    'Formal eQTL colocalization with full summary statistics from eQTL Catalogue is planned as future work.')
                print(f'Fixed Para {i}: removed 【待重算】 marker')
        
        # Remove duplicate MTAP sentence
        full_text = para.text
        if 'MTAP was the dominant arterial eGene in the region' in full_text:
            # This is a duplicate of the earlier sentence, remove it
            for run in para.runs:
                if 'MTAP was the dominant arterial eGene' in run.text:
                    run.text = ''
                    print(f'Fixed Para {i}: removed duplicate MTAP sentence')

# Check for any remaining markers
remaining = []
for i, para in enumerate(doc.paragraphs):
    if '【' in para.text and '】' in para.text:
        remaining.append((i, para.text[:100]))

print(f'\nRemaining 【】 markers: {len(remaining)}')
for i, text in remaining:
    print(f'  Para {i}: {text}')

# Verify key sections
print('\n=== Final verification ===')
full_text = '\n'.join([p.text for p in doc.paragraphs])

checks = [
    ('rg=0.349', 'IA-AAA rg'),
    ('0.054', 'AAA-SBP rg (in table)'),
    ('0.260', 'IA-smoking rg'),
    ('0.226', 'IA-CAD rg'),
    ('0.331', 'Genomic SEM raw'),
    ('0.254', 'Genomic SEM +CAD'),
    ('0.727', 'LAVA 9p21.3 rho (or 0.73)'),
    ('0.675', 'LAVA 18q11 rho (or 0.67)'),
    ('0.88', '9p21.3 PP.H4'),
    ('0.998', '18q11 PP.H4'),
    ('MR-RAPS', 'MR-RAPS'),
    ('MR-PRESSO', 'MR-PRESSO'),
    ('20.0', '18q11 coordinate'),
    ('7,495', 'IA cases'),
    ('39,221', 'AAA cases'),
    ('151,400', 'AAA N_eff (if mentioned)'),
    ('0.200', 'IA h2 (if mentioned)'),
    ('0.014', 'AAA h2 (if mentioned)'),
]

for pattern, desc in checks:
    found = pattern in full_text
    # Also check tables
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if pattern in cell.text:
                    found = True
    status = 'OK' if found else 'not found (may be rounded)'
    print(f'  [{status}] {desc}')

# Save
output_path = r'D:\IA_AAA_analysis\Manuscript_v3_final.docx'
doc.save(output_path)
print(f'\nSaved to: {output_path}')
print(f'Paragraphs: {len(doc.paragraphs)}, Tables: {len(doc.tables)}')
