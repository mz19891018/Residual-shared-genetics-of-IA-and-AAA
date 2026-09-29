#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
Final cleanup of Manuscript v3:
1. Remove review comments section (paragraphs 0-20) and Table 0
2. Fix Table 3 N values for CAD, lipids, height
3. Verify all key numbers
"""
from docx import Document
from docx.oxml.ns import qn
import copy

input_path = r'D:\IA_AAA_analysis\Manuscript_v3_revised.docx'
output_path = r'D:\IA_AAA_analysis\Manuscript_v3_final.docx'

doc = Document(input_path)

# ============================================================
# 1. Remove review comments section (paragraphs before "Title Page")
# ============================================================

# Find the "Title Page" heading
title_page_idx = None
for i, para in enumerate(doc.paragraphs):
    if para.text.strip() == 'Title Page':
        title_page_idx = i
        break

print(f'Title Page found at paragraph {title_page_idx}')

# Remove all paragraphs before Title Page
# We need to remove from the XML directly
body = doc.element.body
paragraphs_to_remove = []
for i, para in enumerate(doc.paragraphs):
    if i >= title_page_idx:
        break
    paragraphs_to_remove.append(para._element)

for elem in paragraphs_to_remove:
    body.remove(elem)

print(f'Removed {len(paragraphs_to_remove)} review comment paragraphs')

# ============================================================
# 2. Remove Table 0 (review comments table)
# ============================================================

# After removing paragraphs, the first table should still be the review table
# Check first table
first_table = doc.tables[0]
first_header = first_table.rows[0].cells[0].text.strip()
print(f'First table header: "{first_header}"')

if '项目' in first_header or 'v7' in first_header:
    # Remove this table
    table_elem = first_table._element
    body.remove(table_elem)
    print('Removed review comments table (Table 0)')
else:
    print('WARNING: First table is not the review table, checking...')

# ============================================================
# 3. Fix Table 3 (now Table 2 after removal) N values
# ============================================================

# After removal, tables are:
# Table 0: LDSC (was Table 1)
# Table 1: Genomic SEM (was Table 2)
# Table 2: Datasets (was Table 3)
# Table 3: Overlapping loci (was Table 4)
# Table 4: MR (was Table 5)

dataset_table = doc.tables[2]  # Was Table 3
print(f'\nDataset table header: {[c.text.strip() for c in dataset_table.rows[0].cells]}')

# Fix each row's N value
n_fixes = {
    'Coronary artery disease': '~1.1M (EUR)',
    'LDL-C, HDL-C, triglycerides': '~400K (EUR)',
    'Height': '~700K (EUR)',
}

for row in dataset_table.rows[1:]:  # Skip header
    trait = row.cells[0].text.strip()
    n_cell = row.cells[2]
    for key, value in n_fixes.items():
        if key in trait:
            # Clear and set new value
            for para in n_cell.paragraphs:
                for run in para.runs:
                    run.text = ''
            if n_cell.paragraphs and n_cell.paragraphs[0].runs:
                n_cell.paragraphs[0].runs[0].text = value
            else:
                n_cell.paragraphs[0].add_run(value)
            print(f'  Fixed {trait}: N = {value}')

# ============================================================
# 4. Verify remaining tables
# ============================================================

print(f'\n=== Remaining tables: {len(doc.tables)} ===')
for ti, table in enumerate(doc.tables):
    header = [c.text.strip()[:25] for c in table.rows[0].cells]
    print(f'  Table {ti}: {header} ({len(table.rows)} rows)')

# ============================================================
# 5. Verify key numbers
# ============================================================

print('\n=== Key number verification ===')
full_text = '\n'.join([p.text for p in doc.paragraphs])

checks = [
    ('rg=0.349', 'IA-AAA rg'),
    ('0.054', 'AAA-SBP rg'),
    ('0.260', 'IA-smoking rg'),
    ('0.226', 'IA-CAD rg'),
    ('0.331', 'Genomic SEM raw'),
    ('0.260', 'Genomic SEM conditioned'),
    ('0.727', 'LAVA 9p21.3 rho'),
    ('0.675', 'LAVA 18q11 rho'),
    ('0.88', '9p21.3 PP.H4'),
    ('0.998', '18q11 PP.H4'),
    ('MR-RAPS', 'MR-RAPS method'),
    ('MR-PRESSO', 'MR-PRESSO method'),
    ('20.0', '18q11 coordinate'),
    ('7,495', 'IA cases'),
    ('39,221', 'AAA cases'),
]

for pattern, desc in checks:
    found = pattern in full_text
    status = 'OK' if found else 'MISSING'
    print(f'  [{status}] {desc}: "{pattern}"')

# Check no remaining 【待核对】 markers
remaining = []
for para in doc.paragraphs:
    if '【待核对' in para.text or '【补充' in para.text or '【删除' in para.text:
        remaining.append(para.text[:80])

print(f'\n=== Remaining 【待核对】 markers: {len(remaining)} ===')
for r in remaining:
    print(f'  {r}')

# Save
doc.save(output_path)
print(f'\nSaved final manuscript to: {output_path}')
print(f'Paragraphs: {len(doc.paragraphs)}, Tables: {len(doc.tables)}')
